import io
import shutil
import subprocess

import numpy as np
import pytest
from fastapi.testclient import TestClient
from PIL import Image, PngImagePlugin

from scanner import scan_bytes, scan_text
from scanner.core import REGISTRY
from scanner.web import app

AI_TEXT = (
    "In today's fast-paced world, it is important to note that artificial intelligence plays a crucial role in "
    "shaping the landscape of modern business. Furthermore, organizations must navigate the complexities of digital "
    "transformation to remain competitive. Moreover, leveraging cutting-edge tools can foster innovation and "
    "streamline operations. Additionally, a holistic approach ensures that companies embrace change seamlessly. "
    "In conclusion, AI is a testament to human ingenuity and a beacon of progress for everyone involved."
)
HUMAN_TEXT = (
    "ok so i finally tried that ramen place on 5th lol. honestly? kinda mid. the broth was fine but they forgot my egg "
    "and the guy at the counter was like, rushing everyone out at 9. my sister loved it tho!! she got the spicy one "
    "and wouldn't shut up about it the whole drive home. idk, maybe i'll go back on a weekday when it's quieter... "
    "or not. there's a taco truck right next to it that smelled way better tbh"
)


def png_bytes(arr, text=None):
    buf = io.BytesIO()
    info = PngImagePlugin.PngInfo()
    for k, v in (text or {}).items():
        info.add_text(k, v)
    Image.fromarray(arr).save(buf, "PNG", pnginfo=info)
    return buf.getvalue()


def test_every_media_kind_has_analyzers():
    kinds = {k for spec in REGISTRY for k in spec.kinds}
    assert {"text", "image", "audio", "video", "document"} <= kinds


def test_ai_text_scores_higher_than_casual_text():
    ai, human = scan_text(AI_TEXT)["verdict"], scan_text(HUMAN_TEXT)["verdict"]
    assert ai["ai_probability"] > 0.6 > human["ai_probability"]
    assert ai["band"] != "likely_ai"  # heuristics alone must never claim certainty


def test_short_text_is_skipped_not_guessed():
    r = scan_text("Hello there, how are you?")
    assert r["verdict"]["band"] == "inconclusive"
    assert any("words" in s["reason"] for s in r["skipped"])


def test_hidden_characters_flagged():
    r = scan_text(AI_TEXT.replace(" ", " ", 5))
    assert any(s["name"] == "Invisible characters" for s in r["signals"])


def test_stable_diffusion_parameters_are_verified_ai():
    arr = np.zeros((512, 512, 3), np.uint8)
    data = png_bytes(arr, {"parameters": "a cat\nSteps: 20, Sampler: Euler a, CFG scale: 7, Seed: 1"})
    assert scan_bytes(data, "x.png")["overall"]["band"] == "verified_ai"


def test_periodic_artifacts_detected():
    rng = np.random.default_rng(0)
    from scipy.ndimage import gaussian_filter

    base = gaussian_filter(rng.normal(128, 40, (512, 512)), 3)
    yy, xx = np.mgrid[:512, :512]
    checker = ((yy // 2 + xx // 2) % 2) * 2 - 1
    data = png_bytes(np.clip(base + 1.5 * checker, 0, 255).astype(np.uint8))
    sig = next(s for s in scan_bytes(data, "a.png")["signals"] if s["name"] == "Periodic frequency peaks")
    assert sig["score"] > 0.7


def test_document_parts_are_scanned():
    import docx

    d = docx.Document()
    d.add_paragraph(AI_TEXT)
    img = io.BytesIO(png_bytes(np.zeros((256, 256, 3), np.uint8), {"parameters": "Steps: 20, Sampler: DPM"}))
    d.add_picture(img)
    buf = io.BytesIO()
    d.save(buf)
    r = scan_bytes(buf.getvalue(), "essay.docx")
    assert {p["kind"] for p in r["parts"]} == {"text", "image"}
    assert r["overall"]["band"] == "verified_ai"


@pytest.mark.skipif(not shutil.which("ffmpeg"), reason="ffmpeg missing")
def test_band_limited_audio_flagged(tmp_path):
    wav = tmp_path / "a.wav"
    subprocess.run(["ffmpeg", "-v", "quiet", "-f", "lavfi", "-i", "anoisesrc=d=3:c=white:a=0.3:r=22050",
                    "-ar", "44100", str(wav)], check=True)
    r = scan_bytes(wav.read_bytes(), "a.wav")
    assert any(s["name"] == "Hard cutoff at a generator rate" for s in r["signals"])


@pytest.mark.skipif(not shutil.which("ffmpeg"), reason="ffmpeg missing")
def test_video_is_split_into_frames_and_audio(tmp_path):
    mp4 = tmp_path / "v.mp4"
    subprocess.run(["ffmpeg", "-v", "quiet", "-f", "lavfi", "-i", "testsrc2=size=320x240:rate=24:duration=3",
                    "-f", "lavfi", "-i", "sine=d=3", "-shortest", "-metadata", "comment=Generated with Sora", str(mp4)],
                   check=True)
    r = scan_bytes(mp4.read_bytes(), "v.mp4")
    kinds = [p["kind"] for p in r["parts"]]
    assert "image" in kinds and "audio" in kinds
    assert any(s["name"] == "AI tool named in metadata" for s in r["signals"])


def test_web_api():
    c = TestClient(app)
    assert c.get("/").status_code == 200
    r = c.post("/api/scan", data={"text": AI_TEXT})
    assert r.status_code == 200 and r.json()["kind"] == "text"
    r = c.post("/api/scan", files={"file": ("x.png", png_bytes(np.zeros((128, 128, 3), np.uint8)), "image/png")})
    assert r.status_code == 200 and r.json()["kind"] == "image"
    assert c.post("/api/scan").status_code == 400
